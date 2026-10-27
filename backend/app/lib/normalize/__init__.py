"""把各种来源的原始内容归一化为 Markdown 中间格式。

提取层（模式 A 的 AI 直提、模式 B 的脚本）只面对这一种格式，
所以所有格式分支都在这里收敛，别处不要再写解析代码。

图片文件不在这里处理：图片题走多模态直读（视觉模型），见 app.lib.ai。
网页里嵌的题目图片例外——那种页面正文就是图，会在这里下载好放进
Normalized.images，由 sources 落盘、提取层走同一个视觉通道。
"""

import io
import json
import re
from dataclasses import dataclass, field
from urllib.parse import urljoin, urlparse

from app.lib.normalize.html2md import html_to_markdown

MAX_FILE_BYTES = 20 * 1024 * 1024
MAX_TEXT_BYTES = 2 * 1024 * 1024
MAX_URL_BYTES = 10 * 1024 * 1024

# 网页图片：题目页经常整页是图，文字提取出来是空的
MAX_URL_IMAGES = 40                       # 单个网页最多下载多少张
MAX_URL_IMAGE_BYTES = 8 * 1024 * 1024     # 单张上限
MAX_URL_IMAGE_TOTAL = 16 * 1024 * 1024    # 一个网页的图片总预算
MIN_IMAGE_BYTES = 1024                    # 小于这个多半是图标 / 追踪像素

IMAGE_MIMES = {"image/png", "image/jpeg", "image/webp", "image/gif", "image/bmp"}
_TRACKING_HINTS = ("pixel", "spacer", "1x1", "tracking", "beacon", "logo", "avatar", "blank")

TEXT_EXTS = {".md", ".markdown", ".txt"}
CSV_EXTS = {".csv", ".tsv"}
HTML_EXTS = {".html", ".htm"}
PDF_EXTS = {".pdf"}
DOCX_EXTS = {".docx"}
IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp"}


class NormalizeError(Exception):
    """归一化失败，message 是给用户看的中文说明。"""


@dataclass
class Normalized:
    markdown: str
    kind: str                      # document / image / url / text
    meta: dict = field(default_factory=dict)
    truncated: bool = False
    # 网页里下载的题目图片（字节, mime）。只由 normalize_url 产出，
    # 不进 meta（JSON 存不了字节），由 sources 落盘后写进 meta.images。
    images: list[tuple[bytes, str]] = field(default_factory=list)


def _ext(filename: str) -> str:
    return "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""


def _decode(data: bytes) -> str:
    for encoding in ("utf-8", "utf-8-sig", "gb18030", "big5"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise NormalizeError("无法识别文件编码（试过 UTF-8 / GB18030 / Big5）")


def _tidy(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _cap(text: str, limit: int) -> tuple[str, bool]:
    if len(text) <= limit:
        return text, False
    return text[:limit], True


def _from_text_like(filename: str, data: bytes) -> Normalized:
    ext = _ext(filename)
    raw = _decode(data)

    if ext in CSV_EXTS:
        return _from_csv(raw, delimiter="\t" if ext == ".tsv" else ",")

    if ext == ".json":
        try:
            parsed = json.loads(raw)
        except json.JSONDecodeError as error:
            raise NormalizeError(f"JSON 解析失败：{error.msg}（第 {error.lineno} 行）") from error
        pretty = json.dumps(parsed, ensure_ascii=False, indent=2)
        return Normalized(f"```json\n{pretty}\n```", "document", {"ext": ext})

    markdown, truncated = _cap(_tidy(raw), 400_000)
    return Normalized(markdown, "document", {"ext": ext or ".txt", "truncated": truncated})


def _from_csv(raw: str, delimiter: str) -> Normalized:
    import csv

    reader = csv.reader(io.StringIO(raw), delimiter=delimiter)
    rows = [[cell.strip() for cell in row] for row in reader]
    rows = [row for row in rows if any(cell for cell in row)]
    if not rows:
        raise NormalizeError("CSV 里没有数据行")

    width = max(len(row) for row in rows)
    rows = [row + [""] * (width - len(row)) for row in rows]
    lines = ["| " + " | ".join(rows[0]) + " |", "| " + " | ".join(["---"] * width) + " |"]
    lines += ["| " + " | ".join(row) + " |" for row in rows[1:]]
    return Normalized("\n".join(lines), "document", {"rows": len(rows)})


def _mime_from_url(url: str) -> str:
    """按 URL 路径的扩展名猜 mime，猜不出给空串。"""
    path = urlparse(url).path.lower()
    for ext, mime in (
        (".png", "image/png"), (".jpg", "image/jpeg"), (".jpeg", "image/jpeg"),
        (".webp", "image/webp"), (".gif", "image/gif"), (".bmp", "image/bmp"),
    ):
        if path.endswith(ext):
            return mime
    return ""


def _collect_image_urls(html: str, base_url: str) -> list[str]:
    """从 HTML 里挑出可能是题目内容的图片 URL。

    相对路径按 base_url 补全；data: 内联图、svg / ico、名字像追踪像素的
    一律跳过。懒加载站点的图在 data-src 里，优先看它。
    """
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(html, "html.parser")
    seen: set[str] = set()
    urls: list[str] = []

    for img in soup.find_all("img"):
        candidate = ""
        for attr in ("data-src", "data-original", "data-actualsrc", "src"):
            value = (img.get(attr) or "").strip()
            if value:
                candidate = value
                break
        if not candidate:
            srcset = (img.get("srcset") or "").strip()
            if srcset:
                candidate = srcset.split(",")[0].strip().split(" ")[0]
        if not candidate or candidate.startswith("data:"):
            continue

        absolute = urljoin(base_url, candidate)
        parsed = urlparse(absolute)
        if parsed.scheme not in ("http", "https"):
            continue
        path = parsed.path.lower()
        if path.endswith((".svg", ".ico")):
            continue
        if any(hint in path for hint in _TRACKING_HINTS):
            continue
        if absolute in seen:
            continue
        seen.add(absolute)
        urls.append(absolute)
        if len(urls) >= MAX_URL_IMAGES:
            break

    return urls


def _fetch_page_images(html: str, base_url: str) -> list[tuple[bytes, str]]:
    """把网页里的题目图片下载下来。失败的单张跳过，不影响整体。"""
    import httpx

    urls = _collect_image_urls(html, base_url)
    if not urls:
        return []

    images: list[tuple[bytes, str]] = []
    total = 0
    try:
        with httpx.Client(
            follow_redirects=True,
            timeout=httpx.Timeout(15.0, connect=8.0),
            max_redirects=3,
            headers={"User-Agent": "Mozilla/5.0 (compatible; MojuanImporter/1.0)"},
        ) as client:
            for url in urls:
                if len(images) >= MAX_URL_IMAGES or total >= MAX_URL_IMAGE_TOTAL:
                    break
                try:
                    response = client.get(url)
                except httpx.HTTPError:
                    continue
                if response.status_code >= 400:
                    continue
                mime = response.headers.get("content-type", "").split(";")[0].strip().lower()
                if mime not in IMAGE_MIMES:
                    # 有的服务器不给 content-type，按扩展名兜底
                    mime = _mime_from_url(url)
                if mime not in IMAGE_MIMES:
                    continue
                declared = response.headers.get("content-length", "")
                if declared.isdigit() and int(declared) > MAX_URL_IMAGE_BYTES:
                    continue
                data = response.content[:MAX_URL_IMAGE_BYTES]
                if len(data) < MIN_IMAGE_BYTES:
                    continue
                images.append((data, mime))
                total += len(data)
    except httpx.HTTPError:
        return images
    return images


def _from_html(data: bytes, base_url: str | None = None) -> Normalized:
    raw = _decode(data)
    markdown, truncated = _cap(html_to_markdown(raw), 400_000)
    # 图片型网页：正文是图，文字提取出来是空的。把图下载下来走视觉通道，
    # 不能再以「没文字」为由拒掉。上传的 HTML 文件没有 base_url，不收集。
    images = _fetch_page_images(raw, base_url) if base_url else []
    if not markdown.strip() and not images:
        raise NormalizeError("网页里没提取到文字内容，也没找到可读取的图片（可能是纯 JS 渲染的页面）")
    title = re.search(r"<title[^>]*>(.*?)</title>", raw, re.IGNORECASE | re.DOTALL)
    meta = {"title": title.group(1).strip()[:120]} if title else {}
    meta["truncated"] = truncated
    if images:
        meta["imageCount"] = len(images)
    return Normalized(markdown, "document", meta, images=images)


def _from_pdf(data: bytes) -> Normalized:
    try:
        from pypdf import PdfReader
    except ImportError as error:  # pragma: no cover
        raise NormalizeError("服务端缺少 pypdf，无法解析 PDF") from error

    try:
        reader = PdfReader(io.BytesIO(data))
    except Exception as error:
        raise NormalizeError(f"PDF 打开失败：{error}") from error

    pages: list[str] = []
    for index, page in enumerate(reader.pages):
        try:
            text = page.extract_text() or ""
        except Exception:
            text = ""
        if text.strip():
            pages.append(f"<!-- 第 {index + 1} 页 -->\n{_tidy(text)}")

    if not pages:
        raise NormalizeError("PDF 里提取不到文字（可能是扫描件，需要 OCR，暂不支持）")

    markdown, truncated = _cap("\n\n".join(pages), 400_000)
    return Normalized(markdown, "document", {"pages": len(reader.pages), "truncated": truncated})


def _from_docx(data: bytes) -> Normalized:
    try:
        from docx import Document
    except ImportError as error:  # pragma: no cover
        raise NormalizeError("服务端缺少 python-docx，无法解析 Word") from error

    try:
        document = Document(io.BytesIO(data))
    except Exception as error:
        raise NormalizeError(f"Word 打开失败：{error}") from error

    # python-docx 把段落与表格分成两条独立列表，拿不到真正的文档流顺序。
    # 对题目材料来说「先全部段落、再全部表格」够用，不引 lxml 去拼 body 顺序。
    lines: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style = (paragraph.style.name or "").lower()
        if style.startswith("heading"):
            level = style.replace("heading", "").strip() or "1"
            level = level if level.isdigit() and 1 <= int(level) <= 6 else "1"
            lines.append(f"{'#' * int(level)} {text}")
        elif style in ("title",):
            lines.append(f"# {text}")
        else:
            lines.append(text)

    for table in document.tables:
        rows = [[" ".join(cell.text.split()) for cell in row.cells] for row in table.rows]
        rows = [row for row in rows if any(cell for cell in row)]
        if not rows:
            continue
        width = max(len(row) for row in rows)
        rows = [row + [""] * (width - len(row)) for row in rows]
        block = ["| " + " | ".join(rows[0]) + " |", "| " + " | ".join(["---"] * width) + " |"]
        block += ["| " + " | ".join(row) + " |" for row in rows[1:]]
        lines.append("\n".join(block))

    markdown, truncated = _cap(_tidy("\n\n".join(lines)), 400_000)
    if not markdown.strip():
        raise NormalizeError("Word 文档里没有文字内容")
    return Normalized(markdown, "document", {"paragraphs": len(lines), "truncated": truncated})


def normalize_bytes(filename: str, data: bytes) -> Normalized:
    """按扩展名分发。图片单独返回，交给视觉模型，不在这里转文本。"""
    if len(data) > MAX_FILE_BYTES:
        raise NormalizeError(f"文件超过 {MAX_FILE_BYTES // 1024 // 1024} MB 上限")

    ext = _ext(filename)
    if ext in IMAGE_EXTS:
        return Normalized("", "image", {"ext": ext, "bytes": len(data), "needsVision": True})

    if ext in PDF_EXTS:
        return _from_pdf(data)
    if ext in DOCX_EXTS:
        return _from_docx(data)
    if ext in HTML_EXTS:
        return _from_html(data)
    if ext in TEXT_EXTS or ext in CSV_EXTS or ext == ".json":
        return _from_text_like(filename, data)

    raise NormalizeError(
        f"暂不支持 {ext or '该'} 格式。目前支持：pdf docx html md txt csv json，以及常见图片"
    )


def normalize_pasted(text: str) -> Normalized:
    """用户直接粘贴的原文：当作纯文本，但保留其中的 Markdown / HTML 结构。"""
    text = text.strip()
    if not text:
        raise NormalizeError("粘贴的内容是空的")
    if len(text.encode("utf-8")) > MAX_TEXT_BYTES:
        raise NormalizeError(f"粘贴内容超过 {MAX_TEXT_BYTES // 1024 // 1024} MB 上限")

    stripped = text.lstrip()
    if stripped[:15].lower().startswith("<!doctype") or stripped[:6].lower() == "<html":
        markdown = html_to_markdown(text)
        kind = "document"
    else:
        markdown = _tidy(text)
        kind = "text"
    markdown, truncated = _cap(markdown, 400_000)
    return Normalized(markdown, kind, {"chars": len(markdown), "truncated": truncated})


def normalize_url(url: str) -> Normalized:
    """抓取网页并归一化。只取文字，不执行 JS。

    网页里嵌的题目图片会一并下载（走视觉模型直读）；
    链接直接指向图片时，整份当作图片材料。
    """
    import httpx

    if not re.match(r"^https?://", url, re.IGNORECASE):
        raise NormalizeError("链接必须以 http:// 或 https:// 开头")

    try:
        with httpx.Client(
            follow_redirects=True,
            timeout=httpx.Timeout(20.0, connect=10.0),
            max_redirects=5,
            headers={"User-Agent": "Mozilla/5.0 (compatible; MojuanImporter/1.0)"},
        ) as client:
            response = client.get(url)
    except httpx.HTTPError as error:
        raise NormalizeError(f"抓取失败：{error}") from error

    if response.status_code >= 400:
        raise NormalizeError(f"目标返回 {response.status_code}")

    content_type = response.headers.get("content-type", "").lower()
    body = response.content[:MAX_URL_BYTES]

    if "pdf" in content_type or url.lower().split("?")[0].endswith(".pdf"):
        return _from_pdf(body)

    # 链接直接指向图片：按图片材料登记，字节交给视觉模型。
    # 有的服务器不给 content-type，按 URL 扩展名兜底。
    mime = content_type.split(";")[0].strip()
    if mime not in IMAGE_MIMES:
        mime = _mime_from_url(url)
    if mime in IMAGE_MIMES:
        ext = ".jpg" if mime == "image/jpeg" else mime.split("/")[-1]
        return Normalized(
            "",
            "image",
            {"ext": ext, "bytes": len(body), "mime": mime, "needsVision": True, "fromUrl": True},
            images=[(body, mime)],
        )

    if any(marker in content_type for marker in ("html", "xml")) or "text/" in content_type:
        # response.url 是重定向后的最终地址，相对图片路径按它补全
        return _from_html(body, base_url=str(response.url))
    return _from_text_like(url.split("/")[-1] or "page.txt", body)
